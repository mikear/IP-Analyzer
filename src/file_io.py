import csv
import io
import json
import logging
from pathlib import Path
from typing import Dict, List, Any, Union

try:
    from docx import Document
    _docx_available = True
except ImportError:
    _docx_available = False

try:
    from fpdf import FPDF
    _fpdf_available = True
except ImportError:
    _fpdf_available = False

logger = logging.getLogger(__name__)

def read_input_file(filepath: Path) -> Union[str, None]:
    """Lee contenido de TXT, CSV, LOG o DOCX."""
    try:
        suffix = filepath.suffix.lower()
        logger.info(f"Intentando leer archivo: {filepath.name} (formato {suffix})")
        if suffix in [".txt", ".csv", ".log"]:
            try:
                content = filepath.read_text(encoding="utf-8")
                logger.info("Archivo leído con éxito (UTF-8).")
                return content
            except UnicodeDecodeError:
                logger.warning(f"Falló lectura UTF-8 de {filepath.name}. Intentando latin-1.")
                try:
                    content = filepath.read_text(encoding="latin-1", errors="replace")
                    logger.info("Archivo leído con éxito (Latin-1 con reemplazo).")
                    return content
                except Exception as latin_err:
                    logger.error(f"Falló también la lectura con Latin-1: {latin_err}")
                    return None
            except Exception as txt_err:
                logger.error(f"Error leyendo archivo de texto {filepath.name}: {txt_err}")
                return None
        elif suffix == ".docx":
            if not _docx_available:
                logger.critical("Falta 'python-docx' para leer archivos .docx.")
                logger.critical("Instala con: pip install python-docx")
                return None
            try:
                doc = Document(filepath)
                full_text = "\n".join(
                    [p.text for p in doc.paragraphs if p.text and p.text.strip()]
                )
                if not full_text.strip():
                    logger.warning(f"El archivo .docx '{filepath.name}' parece vacío.")
                else:
                    logger.info("Archivo .docx leído con éxito.")
                return full_text
            except Exception as docx_err:
                logger.error(
                    f"Error específico leyendo .docx {filepath.name}: {docx_err}",
                    exc_info=True,
                )
                return None
        else:
            logger.error(f"Formato de archivo no soportado: '{suffix}'")
            return None
    except FileNotFoundError:
        logger.critical(f"Archivo no encontrado en la ruta: {filepath}")
        return None
    except IOError as io_err:
        logger.error(f"Error de E/S leyendo el archivo {filepath}: {io_err}")
        return None
    except Exception as e:
        logger.error(f"Error inesperado leyendo el archivo {filepath}: {e}", exc_info=True)
        return None

def _prepare_export_data(results: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Prepara datos para exportación/visualización (aplana, combina)."""
    export_list = []
    if not results: return []
    for idx, result in enumerate(results, start=1):
        ip_address = result.get('ip_address', 'N/A')
        raw_timestamp_str = result.get('raw_timestamp_str', 'N/A')
        timestamp_utc = result.get('original_timestamp_utc_str', 'N/A')
        timestamp_converted = result.get('converted_timestamp', 'N/A')
        ip_info = result.get('ip_info', {})

        flat_data = {
            "orden": idx, "ip_address": ip_address,
            "raw_timestamp_str": raw_timestamp_str,
            "timestamp_utc": timestamp_utc,
            "timestamp_converted": timestamp_converted,
        }
        ip_error = ip_info.get('error')
        if ip_error:
            flat_data.update({
                "isp": f"Error: {ip_error}", "city": "N/A", "region": "N/A",
                "country": "N/A", "hostname": "N/A", "location": "N/A",
                "ip_info_error": ip_error
            })
        else:
            city = ip_info.get('city') or "N/A"
            region = ip_info.get('region') or "N/A"
            country = ip_info.get('country') or "N/A"
            hostname = ip_info.get('hostname') or "N/A"
            isp = ip_info.get('isp') or "N/A"
            location_parts = [p for p in [city, region, country] if p and p != 'N/A']
            location_str = ", ".join(location_parts) if location_parts else "N/A"
            flat_data.update({
                "isp": isp, "city": city, "region": region, "country": country,
                "hostname": hostname, "location": location_str, "ip_info_error": None
            })
        export_list.append(flat_data)
    return export_list

def format_report(results: List[Dict[str, Any]], requested_timezone: str, metadata: Dict[str, str]) -> str:
    """Genera un informe de texto plano formateado."""
    output = io.StringIO(); report_width = 180
    print("=" * report_width, file=output)
    print(f"{ 'INFORME DE ANÁLISIS DE IPs Y ISPs':^{report_width}}", file=output)
    print("=" * report_width, file=output)

    if metadata:
        print("\n--- Datos del Caso ---", file=output)
        # Add SHA256 and App Version to metadata for display
        if "input_file_sha256" in metadata and metadata["input_file_sha256"]:
            print(f"SHA256 del Archivo de Entrada: {metadata['input_file_sha256']}", file=output)
        if "app_version" in metadata and metadata["app_version"]:
            print(f"Versión de la Aplicación: {metadata['app_version']}", file=output)

        filtered_meta = {k: v for k, v in metadata.items() if v and k not in ["input_file_sha256", "app_version", "analysis_start_time", "analysis_duration_seconds", "input_filepath", "target_timezone"]}
        if filtered_meta:
            max_key_len = max(len(key.replace('_', ' ').title()) for key in filtered_meta) + 1
            for key, value in filtered_meta.items(): print(f"{key.replace('_', ' ').title()}:".ljust(max_key_len) + f" {value}", file=output)
        else: print("  (No se proporcionaron datos del caso adicionales)", file=output)
        print("-" * report_width, file=output)

    final_tz_used = requested_timezone
    print(f"\n{ 'Resultados (Zona Horaria Aplicada: ' + final_tz_used + ')':^{report_width}}", file=output)
    print("-" * report_width, file=output)

    if not results:
        print("\nNo se encontraron o procesaron datos válidos.".center(report_width), file=output)
        print("\n" + "=" * report_width, file=output)
        return output.getvalue()

    prepared_data = _prepare_export_data(results)
    headers = {"orden": "Nº", "ip_address": "IP Address", "timestamp_utc": "Timestamp (UTC)", "timestamp_converted": f"Timestamp ({final_tz_used})", "isp": "ISP / Error", "location": "Ubicación", "hostname": "Hostname"}
    widths = {"orden": 4, "ip_address": 38, "timestamp_utc": 23, "timestamp_converted": 28, "isp": 30, "location": 30, "hostname": 24}
    header_keys = list(headers.keys())
    separator = " | "; total_w = sum(widths[k] for k in header_keys) + (len(headers) - 1) * len(separator)
    if total_w > report_width: report_width = total_w
    logger.debug(f"(Ancho tabla TXT: {total_w} caracteres)")

    header_line = separator.join([f"{headers[h]:<{widths[h]}}" for h in header_keys])
    print(header_line, file=output)
    print("-" * total_w, file=output)

    for row_dict in prepared_data:
        row_values = []
        for h_key in header_keys:
            value = str(row_dict.get(h_key, 'N/A')); width = widths[h_key]
            truncated = value[:width-3] + "..." if len(value) > width else value
            align = '^' if h_key == 'orden' else '<'
            row_values.append(f"{truncated:{align}{width}}")
        print(separator.join(row_values), file=output)

    print("\n" + "=" * report_width, file=output)
    print(f"{ 'Fin del informe.':^{report_width}}", file=output)
    print("=" * report_width, file=output)
    return output.getvalue()

def export_to_csv(filepath: Union[str, Path], results: List[Dict[str, Any]], metadata: Dict[str, str]) -> None:
    filepath = Path(filepath); prepared_data = _prepare_export_data(results)
    if not prepared_data: logger.warning("No hay datos para exportar a CSV."); return
    fieldnames = list(prepared_data[0].keys())
    if "ip_info_error" in fieldnames: fieldnames.remove("ip_info_error"); fieldnames.append("ip_info_error")
    try:
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as csvfile:
            if metadata:
                csvfile.write("# --- Metadatos ---\n")
                if "input_file_sha256" in metadata and metadata["input_file_sha256"]:
                    csvfile.write(f"# SHA256 del Archivo de Entrada: {metadata['input_file_sha256']}\n")
                if "app_version" in metadata and metadata["app_version"]:
                    csvfile.write(f"# Versión de la Aplicación: {metadata['app_version']}\n")
                # Write other metadata
                for k, v in metadata.items():
                    if k not in ["input_file_sha256", "app_version", "analysis_start_time", "analysis_duration_seconds", "input_filepath", "target_timezone"]:
                        csvfile.write(f"# {k.replace('_',' ').title()}: {v}\n")
                csvfile.write("# ---\n\n")
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames, extrasaction='ignore')
            writer.writeheader(); writer.writerows(prepared_data)
        logger.info(f"Informe exportado a CSV: {filepath}")
    except IOError as e: logger.error(f"Error E/S escribiendo CSV {filepath}: {e}"); raise
    except Exception as e: logger.error(f"Error inesperado exportando CSV: {e}", exc_info=True); raise

def export_to_json(filepath: Union[str, Path], results: List[Dict[str, Any]], metadata: Dict[str, str]) -> None:
    filepath = Path(filepath); prepared_data = _prepare_export_data(results)
    export_structure = {"metadata": metadata, "results": prepared_data}
    try:
        with open(filepath, 'w', encoding='utf-8') as f: json.dump(export_structure, f, indent=2, ensure_ascii=False)
        logger.info(f"Informe exportado a JSON: {filepath}")
    except IOError as e: logger.error(f"Error E/S escribiendo JSON {filepath}: {e}"); raise
    except Exception as e: logger.error(f"Error inesperado exportando JSON: {e}", exc_info=True); raise

def export_to_txt(filepath: Union[str, Path], results: List[Dict[str, Any]], metadata: Dict[str, str]) -> None:
    filepath = Path(filepath)
    requested_tz = metadata.get("zona_horaria_solicitada_gui", metadata.get("zona_horaria_cli", "UTC"))
    report_content = format_report(results, requested_tz, metadata)
    try:
        with open(filepath, 'w', encoding='utf-8') as f: f.write(report_content)
        logger.info(f"Informe exportado a TXT: {filepath}")
    except IOError as e: logger.error(f"Error E/S escribiendo TXT {filepath}: {e}"); raise
    except Exception as e: logger.error(f"Error inesperado exportando TXT: {e}", exc_info=True); raise

class IPAnalyzerPDF(FPDF):
    def __init__(self, orientation='L', unit='mm', format='A4', metadata: Dict[str, str] = None):
        super().__init__(orientation, unit, format)
        self.app_metadata = metadata if metadata is not None else {}
        self.set_font("Helvetica", size=8) # Default font for footer

    def footer(self):
        # Position at 1.5 cm from bottom
        self.set_y(-15)
        # Set font
        self.set_font("Helvetica", size=8)

        # Page number (right aligned)
        page_num_text = f"Pág. {self.page_no()}/{{nb}}" # {nb} is a placeholder for total pages
        self.cell(0, 10, page_num_text, 0, 0, 'R')

        # Custom footer text (centered)
        app_name = self.app_metadata.get("app_version", "Aplicación").split(" v")[0] # Extract name from "Name vX.Y"
        app_version = self.app_metadata.get("app_version", "")
        developer_name = "Diego A. Rábalo" # Updated developer name
        linkedin_url = "https://www.linkedin.com/in/rabalo" # Updated LinkedIn URL

        footer_text_part1 = f"Informe creado por {app_name} {app_version}, desarrollado por {developer_name}"
        footer_text_part2 = f" ({linkedin_url})" # LinkedIn URL as text

        # Calculate width of the combined text
        combined_text_width = self.get_string_width(footer_text_part1 + footer_text_part2)
        
        # Calculate X position for centering
        center_x = (self.w - combined_text_width) / 2
        self.set_x(center_x)

        # Print first part of the footer text
        self.set_text_color(0) # Black color for normal text
        self.set_font("Helvetica", size=8) # Normal font
        self.cell(self.get_string_width(footer_text_part1), 10, footer_text_part1, 0, 0, 'L')

        # Print LinkedIn URL as clickable text
        self.set_text_color(0, 0, 255) # Blue color for link
        self.set_font("Helvetica", size=8, style='U') # Underline for link
        self.cell(self.get_string_width(footer_text_part2), 10, footer_text_part2, 0, 0, 'L', link=linkedin_url)
        
        # Reset color and font for subsequent text (if any)
        self.set_text_color(0)
        self.set_font("Helvetica", size=8) # Reset font

def export_to_pdf(filepath: Union[str, Path], results: List[Dict[str, Any]], metadata: Dict[str, str]) -> None:
    filepath = Path(filepath)
    if not _fpdf_available:
        logger.critical("Falta 'fpdf2' para generar PDF. Instala: pip install fpdf2")
        raise ImportError("Dependencia FPDF2 no encontrada para exportar a PDF.")

    # Instantiate custom PDF class, passing metadata
    pdf = IPAnalyzerPDF(orientation='L', unit='mm', format='A4', metadata=metadata)
    pdf.alias_nb_pages() # This is crucial for {nb} to work
    pdf.add_page(); pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=8); page_width = pdf.w - 2 * pdf.l_margin

    pdf.set_font("Helvetica", 'B', size=14)
    pdf.cell(page_width, 10, "Informe de Análisis de IPs y ISPs", ln=True, align='C'); pdf.ln(5)

    if metadata:
        pdf.set_font("Helvetica", 'B', size=9); pdf.cell(page_width, 7, "Datos del Caso:", ln=True)
        pdf.set_font("Helvetica", size=8)
        # Add SHA256 and App Version to metadata for display
        if "input_file_sha256" in metadata and metadata["input_file_sha256"]:
            pdf.multi_cell(page_width, 4.5, f"  SHA256 del Archivo de Entrada: {metadata['input_file_sha256']}", ln=True)
        if "app_version" in metadata and metadata["app_version"]:
            pdf.multi_cell(page_width, 4.5, f"  Versión de la Aplicación: {metadata['app_version']}", ln=True)
        
        # Add Total Pages to metadata for display
        # This will be a placeholder for now, updated after content is added
        pdf.multi_cell(page_width, 4.5, f"  Total páginas: {{nb}}", ln=True) # Placeholder for total pages
        
        # Print other metadata
        for k, v in metadata.items():
            if k not in ["input_file_sha256", "app_version", "analysis_start_time", "analysis_duration_seconds", "input_filepath", "target_timezone"]:
                pdf.multi_cell(page_width, 4.5, f"  {k.replace('_',' ').title()}: {v}", ln=True)
        pdf.ln(4)

    requested_tz = metadata.get("zona_horaria_solicitada_gui", metadata.get("zona_horaria_cli", "UTC"))
    pdf.set_font("Helvetica", size=8)
    pdf.cell(page_width, 5, f"Zona Horaria Aplicada: {requested_tz}", ln=True); pdf.ln(5)

    if not results:
        pdf.set_font("Helvetica", 'I', size=10)
        pdf.cell(page_width, 10, "No se encontraron datos válidos.", ln=True, align='C')
    else:
        pdf_data = _prepare_export_data(results)
        headers = ["Nº", "IP Address", "TS (UTC)", f"TS ({requested_tz})", "ISP/Error", "Ubicación", "Hostname"]
        data_keys = ["orden", "ip_address", "timestamp_utc", "timestamp_converted", "isp", "location", "hostname"]
        col_w = {'orden': 10, 'ip_address': 45, 'timestamp_utc': 45, 'timestamp_converted': 50, 'isp': 45, 'location': 45, 'hostname': 37}
        total_w = sum(col_w.values())
        if total_w > page_width:
             logger.warning(f"PDF anchos ({total_w}mm) exceden página ({page_width}mm). Ajustando...")
             scale = page_width / total_w; col_w = {k: v * scale for k, v in col_w.items()}

        line_height = 4.5 # Altura estimada por línea de texto
        header_height = 6
        pdf.set_font("Helvetica", 'B', size=7); pdf.set_fill_color(230); pdf.set_line_width(0.2); pdf.set_text_color(0)
        for key in data_keys: pdf.cell(col_w[key], header_height, headers[data_keys.index(key)], border=1, align='C', fill=True)
        pdf.ln()

        pdf.set_font("Helvetica", size=7)
        for row_dict in pdf_data:
            # Calcular altura necesaria para la fila
            max_lines = 1
            for key in data_keys:
                 value = str(row_dict.get(key, ''))
                 lines = pdf.multi_cell(col_w[key], line_height, value, border=0, align='L', split_only=True)
                 max_lines = max(max_lines, len(lines))
            row_height = max_lines * line_height

            if pdf.get_y() + row_height > pdf.page_break_trigger:
                 pdf.add_page(orientation='L')
                 pdf.set_font("Helvetica", 'B', size=7); pdf.set_fill_color(230)
                 for key in data_keys: pdf.cell(col_w[key], header_height, headers[data_keys.index(key)], border=1, align='C', fill=True)
                 pdf.ln()
                 pdf.set_font("Helvetica", size=7)

            start_y = pdf.get_y()
            for i, key in enumerate(data_keys):
                value = str(row_dict.get(key, ''))
                align = 'C' if key == 'orden' else 'L'
                current_x = pdf.get_x()
                pdf.multi_cell(col_w[key], line_height, value, border=1, align=align, ln=3, max_line_height=line_height)
                pdf.set_xy(current_x + col_w[key], start_y)
            pdf.ln(row_height)

    try:
        pdf.output(str(filepath))
        logger.info(f"Informe PDF generado: {filepath}")
    except Exception as pdf_err:
        logger.error(f"ERROR generando/guardando PDF {filepath}: {pdf_err}", exc_info=True)
        raise
